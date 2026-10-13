"""Rebuild the Word evidence appendix from verified offline analyses."""
from pathlib import Path
import json, hashlib, datetime as dt
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from report_formulas import equation, sub, sup, frac, sigma

ROOT=Path(__file__).resolve().parents[1]
ENG=ROOT/'results/application-slots-engineering-v1/9e45c8058c6942ef9805ab65cc2f8ab9'
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def read(p):return json.loads(p.read_text(encoding='utf-8-sig'))

def build():
    sources=[ROOT/'results/confirmation-rtt-v1/analysis.json',ROOT/'results/application-slots-confirmation-v1/analysis.json',ENG/'analysis.json']
    rtt,slots,eng=map(read,sources)
    assert all(not j['issues'] for j in [rtt,slots,eng])
    expected={'confirmation-rtt-v1':'741d20c14d7fec43f58615c09ba84e8ae982d85e42d965cf648a7830e2f21f89',
        'application-slots-confirmation-v1':'65c69162a62cfd81c8531b9e9fea8480cd664bcddc4d1c78c4bb2b09e46cda0b'}
    for name,h in expected.items():assert sha(ROOT/'results'/name/'protocol.json')==h
    for s in eng['sources']:assert sha(ENG/s['path'])==s['sha256']
    formal_verified=0
    for block in slots['rows']:
        folder=ROOT/'results/application-slots-confirmation-v1'/f"block-{block['block']:02d}"/block['attempt']
        verified=read(folder/'analysis.json');assert not verified['issues']
        for source in verified['sources']:
            assert sha(folder/source['path'])==source['sha256']
            formal_verified+=1
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.8)
    for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
        st=doc.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
        for border in st.element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    doc.styles['Normal'].font.size=Pt(11)
    doc.styles['Normal'].paragraph_format.space_after=Pt(6)
    doc.styles['Normal'].paragraph_format.line_spacing=1.1
    doc.styles['Title'].font.size=Pt(22);doc.styles['Heading 1'].font.size=Pt(17)
    doc.styles['Heading 2'].font.size=Pt(13)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('NT531  |  Phụ lục bằng chứng  |  ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    doc.core_properties.author='Nhóm đồ án NT531'
    doc.core_properties.title='Bằng chứng và thiết kế xác nhận hiệu năng mạng AWS'
    def text(s):doc.add_paragraph(s)
    def page(title):doc.add_page_break();doc.add_heading(title,1)
    def table(headers,rows,widths):
        t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
        for col,w in zip(t.columns,widths):col.width=Inches(w)
        borders=OxmlElement('w:tblBorders')
        for edge in ['top','left','bottom','right','insideH','insideV']:
            b=OxmlElement('w:'+edge)
            for k,v in dict(val='single',sz='4',color='D9D9D9').items():b.set(qn('w:'+k),v)
            borders.append(b)
        t._tbl.tblPr.append(borders)
        records=[(t.rows[0],headers,True,0)]+[(t.add_row(),row,False,i) for i,row in enumerate(rows)]
        for row,values,header,n in records:
            for i,(cell,v) in enumerate(zip(row.cells,values)):
                cell.width=Inches(widths[i]);cell.text=str(v);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                prop=cell._tc.get_or_add_tcPr();shade=OxmlElement('w:shd')
                shade.set(qn('w:fill'),'E8EEF5' if header else ('F5F7FA' if n%2 else 'FFFFFF'));prop.append(shade)
                margins=OxmlElement('w:tcMar')
                for side in ['top','left','bottom','right']:
                    m=OxmlElement('w:'+side);m.set(qn('w:w'),'85');m.set(qn('w:type'),'dxa');margins.append(m)
                prop.append(margins)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1.05
                    p.alignment=WD_ALIGN_PARAGRAPH.LEFT if i==0 else WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:run.font.size=Pt(10.5);run.bold=header
        repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
        doc.add_paragraph().paragraph_format.space_after=Pt(1)

    now=dt.datetime.now(dt.timezone(dt.timedelta(hours=7)))
    doc.add_paragraph('Bằng chứng và thiết kế xác nhận hiệu năng mạng AWS',style='Title')
    text('Phụ lục đồ án NT531  •  Cập nhật '+now.strftime('%d/%m/%Y %H:%M')+' giờ Việt Nam')
    text('Trong một phiên kỹ thuật, tăng khả năng xử lý đồng thời của worker đã khắc phục phần lớn trễ hạn ở 800 yêu cầu/giây trên cả Peering và TGW. Đợt xác nhận kiểm tra khả năng lặp lại của hiệu ứng này và chênh lệch RTT qua nhiều thời điểm. Phụ lục trình bày bằng chứng và phương pháp; chưa thay thế báo cáo toàn bộ đồ án.')
    doc.add_heading('Câu hỏi và tiến độ xác nhận',2)
    text('RTT: TGW có RTT khi rảnh cao hơn Peering trên 0,1 ms trong đa số block của cùng cohort hay không? Slots: tại 800 rps, tăng 8 lên 16 slots trên mỗi worker có nâng tỷ lệ đúng hạn trên 10 điểm phần trăm, trung bình qua hai mode, trong đa số block hay không?')
    table(['Nghiên cứu','Block hợp lệ','Ngày đã đo','Mục tiêu'],[
        ['RTT',f"{rtt['completed_blocks']}/16",len(rtt['local_dates']),'16 block ≥4 ngày'],
        ['Slots',f"{slots['completed_blocks']}/16",len(slots['local_dates']),'16 block ≥4 ngày']],[1.6,1.4,1.3,2.6])
    text('Mỗi nghiên cứu chỉ công bố kết luận suy diễn khi đủ số block, ngày và tiêu chí chất lượng. Engineering/pilot không nhập vào đợt xác nhận. Trạng thái mới nhất nằm trong analysis.json của dự án; bản Word là snapshot tại thời điểm ghi trên.')
    doc.add_heading('Kiến trúc và đường đo',2)
    text('Bốn VPC, mỗi VPC một EC2 c6i.large tại us-east-1a, MTU 1500. Peering full mesh có 6 kết nối; TGW có 1 gateway và 4 VPC attachment. EIP dùng quản trị; các phép đo dùng private IP.')
    text('Vận chuyển: A gửi tới B hoặc A/C/D cùng gửi tới B. Ứng dụng: A → B → C/D, hai RPC song song. B là gateway ứng dụng; C/D là worker có giới hạn đồng thời, không phải router AWS. Giữ hạ tầng/cohort và chỉ thay route hoặc slots theo thiết kế.')

    page('Kết quả engineering của can thiệp slots')
    text('Một phiên kỹ thuật có 84.000 yêu cầu đo. Mỗi tổ hợp warm-up 5 giây, drain backlog, đo 15 giây. Worker mô phỏng I/O 10 ms, 4096 bytes; deadline 50 ms, timeout 2 giây. Tracing đầy đủ tắt; log hàng đợi bật giống nhau.')
    rows=[]
    for rate in [600,800]:
        for mode in ['peering','tgw']:
            for c in [8,16]:
                r=next(x for x in eng['rows'] if (x['mode'],x['slots'],x['rate_rps'])==(mode,c,rate))
                rows.append([mode.upper(),c,rate,r['scheduled'],f"{100*r['deadline_success_fraction']:.3f}",f"{r['successful_completion_p95_ms']:.3f}"])
    table(['Mạng','Slots','rps','N','Đúng hạn %','P95 ms'],rows,[1.35,.65,.65,1,1.3,1.95])
    text('Tất cả yêu cầu engineering thành công HTTP/JSON, nhưng có yêu cầu trễ hạn. P95 dùng request thành công; đúng hạn chia cho mọi yêu cầu theo lịch, kể cả lỗi/không gửi được.')
    doc.add_paragraph().add_run().add_picture(str(ROOT/'docs/images/application/slots-800.png'),width=Inches(6.65))
    doc.add_paragraph('Hình 1  Can thiệp slots tại 800 yêu cầu mỗi giây',style='Caption')
    text('Tải log cuối gặp AWS OAuth 429. Giữ attempt failed; recovery chỉ đọc log còn thiếu, không đo lại. Route snapshot cuối lấy sau cleanup. Phiên này đóng góp 0 block xác nhận.')

    page('Hàng đợi và thời gian phục vụ')
    text('Tại 800 rps, request ID ghép A với C/D cho thấy chờ giảm khi tăng slots; phục vụ vẫn cùng cỡ 10–11 ms. Các hiệu thời gian dùng một đồng hồ monotonic trên từng worker.')
    rows=[]
    for mode in ['peering','tgw']:
        for c in [8,16]:
            r=next(x for x in eng['rows'] if (x['mode'],x['slots'],x['rate_rps'])==(mode,c,800))
            for w in r['workers']:rows.append([mode.upper(),c,w['node'].upper(),f"{w['queue_wait_p95_ms']:.4f}",f"{w['service_p95_ms']:.3f}"])
    table(['Mạng','Slots','Worker','P95 chờ ms','P95 phục vụ ms'],rows,[1.4,.65,.9,1.9,2.05])
    equation(doc,[sub('W','q,i'),' = ',sub('t','service start,i'),' − ',sub('t','entered,i')],1)
    equation(doc,[sub('S','i'),' = ',sub('t','service end,i'),' − ',sub('t','service start,i')],2)
    text('i là lời gọi worker; Wq là chờ, S là phục vụ. Mốc lưu theo giây, nhân 1000 để đổi ms. Handler hoàn tất không chứng minh A nhận; không cộng hai P95 worker thành P95 end-to-end.')
    equation(doc,[sub('ρ','nom'),' = ',frac('λs','c')],3)
    text('λ là yêu cầu/giây vào mỗi worker; s = 0,010 giây là delay cấu hình; c là slots. 800 rps cho ρnom = 1 ở 8 slots, 0,5 ở 16 slots. Đây là tải danh nghĩa, không phải công suất AWS đã đo. Arrival đều theo lịch, không giả định Poisson/M/M/1; backlog tăng chưa ở trạng thái dừng.')
    text('ENA allowance delta bằng 0 ở bốn node/tám tổ hợp chỉ nói không quan sát counter này tăng trong cửa sổ ghi. CPU/ENA là chẩn đoán endpoint; không thay queue logs hoặc loại trừ mọi giới hạn/mất gói.')

    page('Kiểm định xác nhận và các ký hiệu')
    equation(doc,[sub('Δ','j'),' = ',sub('RTT̄','TGW,j'),' − ',sub('RTT̄','Peering,j')],4)
    # Word ignores a combining overbar in a plain math run. Use native accents.
    for token in doc.paragraphs[-1]._p.xpath('.//m:t'):
        if token.text == 'RTT̄':
            run=token.getparent();parent=run.getparent();position=parent.index(run)
            token.text='RTT'
            accent=OxmlElement('m:acc');prop=OxmlElement('m:accPr')
            mark=OxmlElement('m:chr');mark.set(qn('m:val'),'̅');prop.append(mark);accent.append(prop)
            base=OxmlElement('m:e');base.append(run);accent.append(base);parent.insert(position,accent)
    text('j là block. Mỗi mode ba lượt ping, 100 gói/lượt. Trong RTT, Zj = 1 nếu Δj > 0,1 ms, ngược lại bằng 0.')
    equation(doc,[sub('p','m,c,j'),' = ',frac(sub('N','đúng hạn'),sub('N','theo lịch'))],5)
    equation(doc,[sub('d','j'),' = ',frac(['(',sub('p','P,16,j'),' − ',sub('p','P,8,j'),') + (',sub('p','T,16,j'),' − ',sub('p','T,8,j'),')'],'2')],6)
    text('P là Peering, T là TGW; c là slots mỗi worker. Công thức (6) dùng 800 rps; với slots, Zj = 1 nếu dj > 0,10. p và d không đơn vị; 0,10 là 10 điểm phần trăm.')
    equation(doc,['K = ',sigma('j = 1','16'),sub('Z','j'),';  θ = P(',sub('Z','j'),' = 1)'],7)
    text('Trong từng nghiên cứu: H₀: θ ≤ 0,5; H₁: θ > 0,5; α = 0,05. Cỡ mẫu suy diễn là block, không phải request/gói tương quan trong phiên.')
    equation(doc,[sub('p','exact'),' = ',sup('2','−16'),sigma('k = K','16'),frac('16!','k! (16 − k)!')],8)
    text('Ngưỡng K ≥ 12/16 có pexact tại K = 12 bằng 0,0384063721. Đây là phép tính thiết kế, không phải p-value thực nghiệm hiện tại. Vẫn thu cả 16 block, ≥4 ngày, start cách nhau ≥2 giờ; slots yêu cầu uptime ≥30 phút.')
    text('Giữ cohort, mã ứng dụng, protocol và thứ tự đã cân bằng. Khoảng cách không tự chứng minh độc lập. Hai kiểm định dùng α = 0,05 riêng, không bảo đảm đồng thời 95% cho tất cả phát biểu.')

    page('Diễn giải và khả năng kiểm chứng')
    doc.add_heading('Kết luận phù hợp với bằng chứng',2)
    text('Engineering hỗ trợ vai trò của giới hạn đồng thời worker trong workload mô phỏng: đổi mạng ở 8 slots/800 rps vẫn để phần lớn yêu cầu trễ hạn; tăng slots trong từng mode làm đúng hạn tăng và chờ giảm. Xác nhận qua các block mới kiểm tra tính lặp lại. Không suy ra 16 slots tối ưu hoặc một dịch vụ mạng thắng mọi workload.')
    text('Khi đủ dữ liệu, báo cáo effect size, K, p-value exact và cận dưới xác suất một phía 95%; slots có CI bootstrap toàn block cho hiệu ứng trung bình phụ trợ. Không đạt kiểm định không chứng minh tương đương. 100% đúng hạn ở một phiên không bảo đảm 100% trong tương lai.')
    doc.add_heading('Liên hệ nghiên cứu và bài giảng',2)
    text('Dean và Barroso đặt vấn đề độ trễ đuôi khi mức sử dụng/độ phức tạp dịch vụ tăng. Đồ án vận dụng câu hỏi vào hai worker nhỏ, không tái tạo quy mô Google. Bài giảng arrivals/queueing liên hệ Wq, S và giả định trạng thái dừng; QoS liên hệ P95/deadline; congestion liên hệ TCP RTT, cwnd, goodput và retransmission.')
    text('ASA nhấn mạnh p-value không phải xác suất giả thuyết đúng và không đo độ lớn hiệu ứng. Kết luận cần phối hợp thiết kế, raw, cơ chế và phạm vi. Collector → Tempo → Grafana đã kiểm tra bằng replay 12 traces/120 spans thật; chưa triển khai hostmetrics EC2 hoặc OTLP live.')
    doc.add_heading('Nguồn và cách tái tạo',2)
    text('Raw và analysis nằm trong results/confirmation-rtt-v1, results/application-slots-confirmation-v1 và results/application-slots-engineering-v1. Script kiểm tra source hash, route, môi trường, từng request và queue logs. Lỗi/recovery được giữ; state, credentials, SSH key và ZIP riêng không công bố.')
    text('[1] Dean J và Barroso LA. The Tail at Scale. Communications of the ACM 56, 2013, 74–80. https://research.google/pubs/the-tail-at-scale/')
    text('[2] ASA. Statement on Statistical Significance and P Values, 2016. https://www.amstat.org/asa/files/pdfs/p-valuestatement.pdf')
    text('[3] Sigman K. Notes on Little’s Law, 2009, phần 1–1.1. https://www.columbia.edu/~ks20/stochastic-I/stochastic-I-LL.pdf')
    text('Công thức và tiêu chí bàn giao đầy đủ: tài liệu 27 và 30 trong thư mục docs. Chương kết quả vận chuyển cũ giữ nguyên dữ liệu lịch sử và không nhập vào cohort xác nhận mới.')
    native_equations=8
    if slots['rows']:
        block=slots['rows'][-1]
        page('Kết quả block xác nhận slots mới nhất')
        text(f"Block {block['block']} bắt đầu UTC {block['started_utc']}. Đã kiểm tra raw, route, môi trường và log theo request ID; phiên này thuộc mẫu xác nhận, không nhập engineering. Tổng tiến độ {slots['completed_blocks']}/16; chưa có kết luận kiểm định hoặc CI liên block.")
        rows=[]
        for rate in [600,800]:
            for mode in ['peering','tgw']:
                for c in [8,16]:
                    r=next(x for x in block['rows'] if (x['mode'],x['slots'],x['rate_rps'])==(mode,c,rate))
                    rows.append([mode.upper(),c,rate,r['scheduled'],f"{100*r['deadline_success_fraction']:.3f}",f"{r['successful_completion_p95_ms']:.3f}"])
        table(['Mạng','Slots','rps','N','Đúng hạn %','P95 ms'],rows,[1.35,.65,.65,1,1.3,1.95])
        text(f"Tăng đúng hạn 8 → 16 slots tại 800 rps, trung bình hai mode, là {float(block['improvement'])*100:.4f} điểm phần trăm. Đây là hiệu ứng của block này; chưa khẳng định xác suất lặp lại trong đa số block. Mẫu số giữ mọi yêu cầu theo lịch; P95 chỉ dùng request thành công.")
        for mode in ['peering','tgw']:
            low=next(r for r in block['rows'] if (r['mode'],r['slots'],r['rate_rps'])==(mode,8,800))
            high=next(r for r in block['rows'] if (r['mode'],r['slots'],r['rate_rps'])==(mode,16,800))
            text(f"{mode.upper()}: P95 chờ C/D ở 8 slots là "+' / '.join(f"{w['queue_wait_p95_ms']:.3f}" for w in low['workers'])+' ms; ở 16 slots là '+' / '.join(f"{w['queue_wait_p95_ms']:.4f}" for w in high['workers'])+' ms. Các số này là quantile từng worker, không cộng thành P95 phía A.')
        text('Mỗi điều kiện dùng fresh server, warm-up và drain; cùng mã ứng dụng/MTU/cohort. SCP chỉ tải sau capture, host key pin qua SSM, đối chiếu size/SHA. Trước block 1 có attempt kỹ thuật bị guard uptime chặn: giữ lỗi và quyết định retry, chưa thu outcome; không chọn lại vì kết quả.')
        queue_source=ROOT/'results/analysis/queue-occupancy-audit.json'
        queue_audit=read(queue_source)
        cohort=next(q for q in queue_audit['study_results'] if q['phase']=='confirmatory_block'
            and (q['block'],q['attempt'])==(block['block'],block['attempt']))
        folder=ROOT/'results/application-slots-confirmation-v1'/f"block-{block['block']:02d}"/block['attempt']
        assert sha(folder/'analysis.json')==cohort['validated_analysis_sha256']
        for condition in cohort['conditions']:
            for source in condition['sources']:assert sha(ROOT/source['path'])==source['sha256']
        sources.append(queue_source)
        page('Số lời gọi đang chờ và định luật Little')
        text(f"Phân tích mô tả từ block slots {block['block']}, không tạo thêm mẫu xác nhận. Ghép request ID đo, loại warm-up và giữ lời gọi hủy. Mỗi worker có cửa sổ riêng chứa trọn từ entry đầu đến exit cuối, gồm đuôi backlog.")
        equation(doc,[sub('λ','C'),' = ',frac('N','T')],9)
        equation(doc,[sub('L','q,C'),' = ',frac([sigma('i = 1','N'),sub('W','q,i')],'T'),
            ' = ',sub('λ','C'),' × ',sub('W','q,mean')],10)
        native_equations=10
        text('N là số lời gọi vào worker; T là độ dài cửa sổ theo giây; Wq,i là thời gian chờ của lời gọi i, giây; Wq,mean là mean của N thời gian chờ. Lq,C là số lời gọi trung bình đang chờ, λC là N/T, lời gọi/s. Không thay mean bằng P95 hoặc N/T bằng rate gửi 800 khi cửa sổ có đuôi.')
        rows=[]
        for mode in ['peering','tgw']:
            for c in [8,16]:
                condition=next(r for r in cohort['conditions'] if (r['mode'],r['slots'],r['configured_arrival_rps'])==(mode,c,800))
                w=next(w for w in condition['workers'] if w['node']=='c')
                rows.append([mode.upper(),c,f"{w['cohort_count_rate_per_second']:.3f}",f"{w['mean_wait_ms']:.6f}",f"{w['mean_queue_occupancy']:.6f}"])
        table(['Mạng','Slots','N/T mỗi s','Mean chờ ms','Mean số chờ C'],rows,[1.25,.65,1.3,1.8,1.9])
        image=ROOT/'docs/images/application/queue-occupancy'/f"block-{block['block']:02d}-queue.png"
        doc.add_paragraph().add_run().add_picture(str(image),width=Inches(6.2))
        doc.add_paragraph('Hình 2  Mean số chờ trong bin 100 ms tại 800 rps',style='Caption')
        text('Hai hàng dùng thang dọc khác nhau; gốc thời gian riêng mỗi worker. Đẳng thức hữu hạn do tích phân hàm chỉ thị bằng tổng khoảng chờ; không cần trạng thái dừng. Kiểm tra bằng event sweep và tổng thời gian cùng raw là nhất quán đại số, không kiểm định độc lập hoặc bằng chứng hàng đợi AWS. Chứng minh và số liệu hai worker ở tài liệu 33.')
    out=ROOT/'docs/31-phu-luc-bang-chung-thong-ke.docx';doc.save(out)
    proof=dict(created_local=now.isoformat(),docx_sha256=sha(out),rtt_blocks=rtt['completed_blocks'],slots_blocks=slots['completed_blocks'],
        sources=[dict(path=p.relative_to(ROOT).as_posix(),sha256=sha(p)) for p in sources],
        engineering_raw_sources_verified=len(eng['sources']),formal_slot_raw_sources_verified=formal_verified,native_equations=native_equations,render_review='pending')
    (ROOT/'results/analysis/confirmation-appendix.json').write_text(json.dumps(proof,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(proof))

if __name__=='__main__':build()
