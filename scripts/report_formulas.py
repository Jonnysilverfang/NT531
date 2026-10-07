"""Explain and calculate report formulas from the existing verified datasets.

Equations are editable Word OMML. The same content is exported to Markdown.
Jain values use each source's own run window, not an aligned common window.
"""
import hashlib
import json
import math
import statistics as st
from pathlib import Path
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

JAIN_URL = 'https://arxiv.org/abs/cs/9809099'
NIST_URL = 'https://www.itl.nist.gov/div898/software/dataplot/refman2/auxillar/coefvari.htm'
MIT_URL = 'https://ocw.mit.edu/courses/6-02-introduction-to-eecs-ii-digital-communication-systems-fall-2012/612ee68db22a5ce5ae171650e8b3d865_9HCUnJB9ovk.pdf'


def fmt(value, digits=4):
    return f'{value:.{digits}f}'.replace('.', ',')


def calculate(summary, root):
    result = {'jain_window_basis': 'Per-source full-run receiver windows; not common-window aligned.',
              'jain_per_run': {}, 'jain_descriptive': {}, 'cv_percent': {}}
    for mode in ('peering', 'tgw'):
        rows = []
        for run in range(1, 6):
            rates = {r['node']: r['receiver_Gbps'] for r in summary['fanin_metrics'][mode] if r['run'] == run}
            assert set(rates) == set('acd') and all(x >= 0 for x in rates.values())
            total = sum(rates.values())
            assert total > 0
            jain = total**2 / (3 * sum(x*x for x in rates.values()))
            assert 1/3 <= jain <= 1 + 1e-12
            rows.append({'run': run, 'receiver_Gbps': rates, 'approx_total_Gbps': total, 'jain': jain})
        result['jain_per_run'][mode] = rows
        values = [r['jain'] for r in rows]
        result['jain_descriptive'][mode] = {'mean_of_run_indices': st.mean(values), 'minimum': min(values), 'maximum': max(values)}
        result['cv_percent'][mode] = {key: summary[key][mode]['sample_sd'] / summary[key][mode]['mean'] * 100
                                     for key in ('rtt', 'tcp_single', 'fanin_approximate_total')}
    example = root / 'results/formal/peering/tcp-single-a-b/A/peering-tcp-single-a-b-20261006T091620Z-uupMz9/run-1.json'
    raw = json.loads(example.read_text(encoding='utf-8-sig'))['end']['sum_received']
    recomputed = 8 * raw['bytes'] / raw['seconds'] / 1e9
    assert math.isclose(recomputed, summary['single_rows']['peering'][0]['ReceiverGbps'], rel_tol=1e-12)
    result['goodput_example'] = {'bytes': raw['bytes'], 'seconds': raw['seconds'], 'receiver_Gbps': recomputed,
                                'source': example.relative_to(root).as_posix(), 'sha256': hashlib.sha256(example.read_bytes()).hexdigest()}
    result['single_peering_vs_tgw_percent'] = (summary['tcp_single']['peering']['mean'] / summary['tcp_single']['tgw']['mean'] - 1) * 100
    result['fanin_tgw_vs_peering_percent'] = (summary['fanin_approximate_total']['tgw']['mean'] / summary['fanin_approximate_total']['peering']['mean'] - 1) * 100
    result['rtt_tgw_minus_peering_ms'] = summary['rtt']['tgw']['mean'] - summary['rtt']['peering']['mean']
    result['limits'] = ['Descriptive statistics from five runs per batch; no inferential test or confidence interval.',
                        'Mean of Jain per-run indices, not Jain of session-average source rates.',
                        'Logical topology counts are not physical link counts or service capacity measurements.',
                        'Physical delay examples are hypothetical, not AWS measurements.']
    return result


# Small math expression tree for native fractions, radicals, limits and indices.
def frac(num, den): return ('frac', num, den)
def sub(base, index): return ('sub', base, index)
def sup(base, power): return ('sup', base, power)
def root(expr): return ('root', expr)
def sigma(index, upper): return ('limits', '∑', index, upper)


def append_math(parent, expr):
    if isinstance(expr, list):
        for part in expr: append_math(parent, part)
    elif isinstance(expr, tuple):
        kind, *args = expr
        tags = {'frac': ('m:f', ('m:num', 'm:den')), 'sub': ('m:sSub', ('m:e', 'm:sub')),
                'sup': ('m:sSup', ('m:e', 'm:sup')), 'limits': ('m:sSubSup', ('m:e', 'm:sub', 'm:sup'))}
        if kind == 'root':
            element = OxmlElement('m:rad'); prop = OxmlElement('m:radPr')
            hidden = OxmlElement('m:degHide'); hidden.set(qn('m:val'), '1'); prop.append(hidden); element.append(prop)
            element.append(OxmlElement('m:deg')); body = OxmlElement('m:e'); append_math(body, args[0]); element.append(body)
        else:
            tag, slots = tags[kind]; element = OxmlElement(tag)
            for slot, value in zip(slots, args):
                body = OxmlElement(slot); append_math(body, value); element.append(body)
        parent.append(element)
    else:
        r = OxmlElement('m:r'); rp = OxmlElement('w:rPr'); fonts = OxmlElement('w:rFonts')
        fonts.set(qn('w:ascii'), 'Cambria Math'); fonts.set(qn('w:hAnsi'), 'Cambria Math'); rp.append(fonts)
        size = OxmlElement('w:sz'); size.set(qn('w:val'), '24'); rp.append(size); r.append(rp)
        t = OxmlElement('m:t'); t.set(qn('xml:space'), 'preserve'); t.text = str(expr); r.append(t); parent.append(r)


def equation(doc, expr, number):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.keep_with_next = True
    op = OxmlElement('m:oMathPara'); body = OxmlElement('m:oMath')
    append_math(body, expr); append_math(body, f'    ({number})'); op.append(body); p._p.append(op)


def pages(summary):
    m = summary['quantitative_methods']
    p = m['jain_per_run']['peering'][0]; t = m['jain_per_run']['tgw'][0]
    g = m['goodput_example']
    # Blocks: text, subheading, editable equation with Markdown LaTeX, table, reference.
    T = lambda text: ('text', text)
    H = lambda text: ('heading', text)
    E = lambda no, expr, latex: ('equation', no, expr, latex)
    R = lambda label, url: ('reference', label, url)
    return [
        ('7 Phương pháp tính và giải thích công thức', [
            T('Phần này quy định ký hiệu, đơn vị và cách tính các chỉ số ở mục 2–5. Ví dụ thực nghiệm lấy từ file đã kiểm tra; ví dụ vật lý giả định được đánh dấu riêng. Các phép tính mô tả một phiên, chưa thay thế phép kiểm định giữa nhiều phiên độc lập.'),
            H('7.1 Trung bình và mức biến động giữa các lượt'),
            E(1, ['x̄ = ', frac([sigma('i = 1', 'n'), sub('x', 'i')], 'n')], r'\bar{x}=\frac{\sum_{i=1}^{n}x_i}{n}'),
            T('Ký hiệu: n là số lượt hợp lệ (n = 5); i là số thứ tự lượt, từ 1 đến n; xᵢ là chỉ số của lượt i; Σ là phép cộng các giá trị theo i; x̄ là trung bình số học. xᵢ và x̄ cùng đơn vị: ms nếu xét RTT, Gbps nếu xét thông lượng. Với RTT, mỗi xᵢ đã là trung bình của 100 gói trong một lượt.'),
            T('Thế số RTT Peering: x̄ = (0,212 + 0,195 + 0,215 + 0,211 + 0,206)/5 = 0,2078 ms. Các lượt đều có 100 phản hồi nên kết quả này cũng bằng trung bình gộp 500 phản hồi ở độ chính xác của CSV. Khi số phản hồi khác nhau, trung bình giữa lượt và trung bình gộp có thể khác nhau.'),
            E(2, ['s = ', root(frac([sigma('i = 1', 'n'), sup(['(', sub('x', 'i'), ' − x̄)'], '2')], 'n − 1'))], r's=\sqrt{\frac{\sum_{i=1}^{n}(x_i-\bar{x})^2}{n-1}}'),
            T('Ký hiệu: s là độ lệch chuẩn mẫu; (xᵢ − x̄) là độ lệch của lượt i so với trung bình; số mũ 2 là bình phương; √ là căn bậc hai. Chia n − 1 để tính phương sai mẫu rồi lấy căn, vì vậy s có cùng đơn vị với xᵢ. Đây là SD giữa các lượt, khác mdev của ping trong từng lượt.'),
            T('Thế số: tổng bình phương độ lệch của năm RTT Peering là 0,0002468 ms²; s = √(0,0002468/4) = 0,0078549 ms. TGW có s = 0,0775081 ms. Mẫu số 4 là n − 1, không phải số gói ICMP.'),
            E(3, ['CV = ', frac('s', 'x̄'), ' × 100%'], r'CV=\frac{s}{\bar{x}}\times100\%'),
            T(f'Ký hiệu: CV là hệ số biến thiên biểu diễn theo phần trăm, dùng khi x̄ > 0 [2]. RTT Peering: 0,0078549/0,2078 × 100% = {fmt(m["cv_percent"]["peering"]["rtt"],2)}%; TGW: {fmt(m["cv_percent"]["tgw"]["rtt"],2)}%. CV nhỏ cho thấy năm trung bình ít biến động tương đối trong phiên; không phải sai số phần trăm, khoảng tin cậy hoặc bằng chứng tái lập qua nhiều ngày.'),
            R('[2] NIST. Coefficient of Variation.', NIST_URL),
        ]),
        ('7.2 RTT mất phản hồi và thông lượng nhận', [
            H('RTT của một gói ICMP'),
            E(4, [sub('r', 'k'), ' = ', sub('t', 'nhận,k'), ' − ', sub('t', 'gửi,k')], r'r_k=t_{\mathrm{nhận},k}-t_{\mathrm{gửi},k}'),
            T('Ký hiệu: k là số thứ tự gói có phản hồi; rₖ là RTT của gói k; t_gửi,k là mốc gửi yêu cầu tại A; t_nhận,k là mốc A nhận phản hồi tương ứng. Hai mốc lấy trên cùng máy A. Nếu thời gian dùng giây, nhân 1000 để đổi RTT sang ms. Ping báo RTT trực tiếp; không cần đồng bộ đồng hồ A và B để tính RTT.'),
            T('Ví dụ thực nghiệm: RTT trung bình của lượt 1 Peering là 0,212 ms. Không coi RTT/2 là độ trễ một chiều đã đo: phép chia đó chỉ là ước lượng khi giả định hai chiều đối xứng và thời gian phản hồi ở B không đáng kể.'),
            H('Tỷ lệ không nhận được phản hồi ICMP'),
            E(5, ['P = ', frac([sub('N', 'gửi'), ' − ', sub('N', 'nhận')], sub('N', 'gửi')), ' × 100%'], r'P=\frac{N_{\mathrm{gửi}}-N_{\mathrm{nhận}}}{N_{\mathrm{gửi}}}\times100\%'),
            T('Ký hiệu: N_gửi là số yêu cầu ICMP đã gửi; N_nhận là số phản hồi nhận được; P có đơn vị %. Điều kiện: N_gửi > 0. Thế số mỗi mode: (500 − 500)/500 × 100% = 0%. Đây là mất phản hồi quan sát bằng ping, không xác định gói bị mất ở chiều đi hay chiều về. Không suy ra mọi lưu lượng TCP cũng không mất gói.'),
            H('Thông lượng phía nhận ở tầng ứng dụng'),
            E(6, ['G = ', frac([ '8 × ', sub('B', 'nhận')], ['T × ', sup('10', '9')])], r'G=\frac{8B_{\mathrm{nhận}}}{T\times10^9}'),
            T('Ký hiệu: B_nhận là số byte dữ liệu nhận được trong cửa sổ báo cáo; T là thời lượng chính cửa sổ đó, tính bằng giây; G có đơn vị Gbps. Hệ số 8 đổi byte thành bit; 10⁹ đổi bit/s thành Gbps theo hệ thập phân. Đây là goodput ứng dụng do iperf3 báo, không phải tốc độ vật lý hoặc tổng byte kể cả header và truyền lại.'),
            T(f'Thế số từ end.sum_received của JSON lượt 1 Peering: B_nhận = {g["bytes"]:,} byte; T = {fmt(g["seconds"],5)} s. G = 8 × {g["bytes"]:,}/({fmt(g["seconds"],5)} × 10⁹) = {fmt(g["receiver_Gbps"],6)} Gbps, khớp summary.csv. Dùng thời lượng trong JSON thay vì tự coi T đúng bằng 30 s; không cộng 5 s khởi động đã bỏ vào mẫu số này.'.replace(f'{g["bytes"]:,}', f'{g["bytes"]:,}'.replace(',', ' '))),
            T('Nguồn ví dụ: bộ TCP một luồng Peering tại A, lượt 1; SHA256 của JSON nằm trong results/analysis/report-metrics.json. B là số byte trong công thức (6), không phải tên máy B.'),
        ]),
        ('7.3 So sánh và tổng thông lượng ba nguồn', [
            H('Chênh lệch có đơn vị và chênh lệch phần trăm'),
            E(7, ['D = ', sub('x̄', 'mới'), ' − ', sub('x̄', 'mốc')], r'D=\bar{x}_{\mathrm{mới}}-\bar{x}_{\mathrm{mốc}}'),
            E(8, ['δ = ', frac([sub('x̄', 'mới'), ' − ', sub('x̄', 'mốc')], sub('x̄', 'mốc')), ' × 100%'], r'\delta=\frac{\bar{x}_{\mathrm{mới}}-\bar{x}_{\mathrm{mốc}}}{\bar{x}_{\mathrm{mốc}}}\times100\%'),
            T('Ký hiệu: x̄_mới là trung bình phương án đang xét; x̄_mốc là trung bình phương án được chọn làm mốc (phải > 0); D giữ đơn vị của chỉ số; δ có đơn vị %. D và δ có dấu: dương nghĩa là giá trị mới cao hơn mốc, không tự động có nghĩa là tốt hơn. RTT thấp thường được ưu tiên, còn thông lượng cao thường được ưu tiên. Mỗi phép so sánh phải ghi rõ mẫu số.'),
            T(f'RTT chọn TGW là mới, Peering là mốc: D = 0,5630 − 0,2078 = {fmt(m["rtt_tgw_minus_peering_ms"])} ms. TCP một luồng chọn Peering là mới, TGW là mốc: δ = (4,78280398 − 4,60826505)/4,60826505 × 100% = +{fmt(m["single_peering_vs_tgw_percent"],2)}%.'),
            T(f'Ba nguồn chọn TGW là mới, Peering là mốc: δ = (11,59051996 − 11,61100879)/11,61100879 × 100% = {fmt(m["fanin_tgw_vs_peering_percent"],4)}%. Đây là chênh lệch quan sát của hai phiên nối tiếp. Không ghép lượt 1 Peering với lượt 1 TGW thành cặp đo cùng thời điểm hoặc coi chênh lệch nhỏ là bằng chứng tương đương thống kê.'),
            H('Tổng xấp xỉ khi ba nguồn cùng gửi tới B'),
            E(9, [sub('S', 'i'), ' ≈ ', sub('G', 'A,i'), ' + ', sub('G', 'C,i'), ' + ', sub('G', 'D,i')], r'S_i\approx G_{A,i}+G_{C,i}+G_{D,i}'),
            T('Ký hiệu: i là lượt đo; G_A,i, G_C,i và G_D,i là receiver goodput từ A, C và D trong lượt i, cùng đơn vị Gbps; Sᵢ là tổng xấp xỉ của ba nguồn. Dấu ≈ nhắc rằng các nguồn có cửa sổ riêng hơi lệch nhau. Trung bình của Sᵢ qua năm lượt tương đương cộng trung bình từng nguồn khi mỗi nguồn đều có đủ năm lượt.'),
            T('Phép tính dùng các trung bình đầy đủ; minh họa làm tròn tám chữ số: Peering 3,95832808 + 3,77990339 + 3,87277732 ≈ 11,61100879 Gbps. Nếu cộng ba số đã làm tròn bốn chữ số, kết quả có thể lệch ở chữ số cuối.'),
            T('Muốn tổng chính xác trên khoảng chung, chọn cùng một cửa sổ [t₀,t₁] cho ba nguồn, lấy tổng byte nhận trong chính khoảng đó rồi chia cho t₁ − t₀ theo (6). Bảng hiện tại chưa thực hiện phép căn chỉnh interval này; không diễn giải Sᵢ là công suất tối đa của TGW hoặc Peering.'),
        ]),
        ('7.4 Mức chia đều thông lượng giữa ba nguồn', [
            E(10, [sub('J', 'i'), ' = ', frac(sup(['(', sub('G', 'A,i'), ' + ', sub('G', 'C,i'), ' + ', sub('G', 'D,i'), ')'], '2'), ['3 × (', sup(sub('G', 'A,i'), '2'), ' + ', sup(sub('G', 'C,i'), '2'), ' + ', sup(sub('G', 'D,i'), '2'), ')'])], r'J_i=\frac{(G_{A,i}+G_{C,i}+G_{D,i})^2}{3(G_{A,i}^2+G_{C,i}^2+G_{D,i}^2)}'),
            T('Jᵢ là chỉ số Jain của lượt i [3]; ba G là thông lượng phía nhận của A/C/D trong cùng lượt, đo bằng cùng đơn vị. Hệ số 3 là số nguồn. Với thông lượng không âm và tổng khác 0, 1/3 ≤ Jᵢ ≤ 1; bằng 1 khi ba nguồn có thông lượng bằng nhau. J không có đơn vị; nếu cả ba bằng 0 thì không tính được.'),
            T(f'Thế số lượt 1 Peering (ba G tính bằng Gbps): J₁ = ({fmt(p["receiver_Gbps"]["a"],6)} + {fmt(p["receiver_Gbps"]["c"],6)} + {fmt(p["receiver_Gbps"]["d"],6)})² / [3 × ({fmt(p["receiver_Gbps"]["a"],6)}² + {fmt(p["receiver_Gbps"]["c"],6)}² + {fmt(p["receiver_Gbps"]["d"],6)}²)] ≈ {fmt(p["jain"],6)}. Đơn vị Gbps² triệt tiêu giữa tử và mẫu. Phép tính dùng số đầy đủ trong JSON trước khi làm tròn.'),
            ('table', ['Lượt', 'J Peering', 'J TGW'], [[str(i+1),fmt(m['jain_per_run']['peering'][i]['jain'],6),fmt(m['jain_per_run']['tgw'][i]['jain'],6)] for i in range(5)], [1.3,2.7,2.7]),
            T(f'Trung bình của năm Jᵢ: Peering {fmt(m["jain_descriptive"]["peering"]["mean_of_run_indices"],6)}; TGW {fmt(m["jain_descriptive"]["tgw"]["mean_of_run_indices"],6)}. Phạm vi từng lượt: Peering {fmt(m["jain_descriptive"]["peering"]["minimum"],6)}–{fmt(m["jain_descriptive"]["peering"]["maximum"],6)}; TGW {fmt(m["jain_descriptive"]["tgw"]["minimum"],6)}–{fmt(m["jain_descriptive"]["tgw"]["maximum"],6)}. Trung bình các Jᵢ khác với lấy Jain của ba thông lượng trung bình cả phiên.'),
            T('Kết quả mô tả mức chia đều trong các cửa sổ đã ghi. Nó không đo tổng công suất, không xác định nguyên nhân chia tải và chưa chứng minh một kiến trúc công bằng hơn nói chung. Các nguồn hơi lệch cửa sổ; cần Jain trên khoảng chung và nhiều phiên nếu muốn kết luận chắc hơn.'),
            R('[3] Jain, Chiu và Hawe. A Quantitative Measure of Fairness and Discrimination for Resource Allocation in Shared Computer Systems. DEC TR-301, 1984; bản lưu arXiv.', JAIN_URL),
            T('Các đầu vào G và J đầy đủ của từng lượt được lưu trong quantitative_methods của report-metrics.json. Phép tính bổ sung dùng dữ liệu sẵn có, không tạo một phiên đo mới.'),
        ]),
        ('7.5 CPU và thay đổi bộ đếm ENA', [
            H('Tỷ lệ CPU đang bận'),
            E(11, [sub('U', 'j'), ' = 100 − ', sub('I', 'j')], r'U_j=100-I_j'),
            T('Ký hiệu: j là số thứ tự mẫu mpstat trong cửa sổ; Iⱼ là %idle trên dòng all; Uⱼ là CPU busy, cùng đơn vị %. all tổng hợp các CPU logic của máy. Trung bình Uⱼ theo các mẫu trong cửa sổ rồi trung bình năm cửa sổ để mô tả phiên. Với khoảng lấy mẫu 1 s, các mẫu được cho trọng số bằng nhau.'),
            T('Ví dụ thực nghiệm TCP một luồng Peering tại B, lượt 1: CPU busy trung bình 17,656%, tương ứng idle trung bình 82,344%; 100 − 82,344 = 17,656%. Bài ba nguồn có CPU busy trung bình năm cửa sổ là 59,18% qua Peering và 56,88% qua TGW.'),
            T('CPU busy là mức bận toàn máy, gồm xử lý hệ thống, ngắt và các tiến trình khác. Không gọi đây là CPU riêng của iperf3. Một lõi hoặc giới hạn mạng vẫn có thể là nút thắt dù CPU all thấp hơn 100%. Peering có 29–30 mẫu mỗi cửa sổ ở B; TGW có 30 mẫu. Các giá trị đã giữ số mẫu trong audit.'),
            H('Delta của một bộ đếm tích lũy'),
            E(12, ['ΔC = ', sub('C', 'sau'), ' − ', sub('C', 'trước')], r'\Delta C=C_{\mathrm{sau}}-C_{\mathrm{trước}}'),
            T('Ký hiệu: C_trước và C_sau là giá trị của cùng một bộ đếm trên cùng interface ens5 ở hai mốc chụp; ΔC là phần tăng trong khoảng chụp, đơn vị đếm gói theo định nghĩa bộ đếm. Điều kiện: driver không reset hoặc bộ đếm không quay vòng giữa hai mốc. Nếu C_sau < C_trước thì phải kiểm tra reset, không diễn giải số âm như giảm mất gói.'),
            T('Ví dụ thực nghiệm toàn phiên ba nguồn tại B: Δbw_in_allowance_exceeded = 1 434 501 qua Peering và 6 608 210 qua TGW. Đây là số tăng đã đối chiếu từ ảnh chụp trước/sau; không phải tổng số gói mạng B nhận được.'),
            T('Theo AWS [1], bw_in và pps allowance phản ánh gói bị xếp hàng hoặc loại bỏ vì vượt giới hạn tương ứng. Chúng không tăng chỉ vì thời gian trôi qua. Không cộng hai delta để tính gói mất, không chia cho số gói ping để tính mất gói TCP. Khoảng chụp tại B gồm cả chờ và nghỉ, nên chưa tách delta riêng cho từng lượt hoặc nguồn.'),
            T('Có thể chuẩn hóa ΔC theo thời lượng chụp để mô tả tốc độ tăng bộ đếm nếu có đủ mốc thời gian. Tuy nhiên tốc độ đó vẫn là bộ đếm allowance, không phải PPS tổng hoặc tỷ lệ mất gói. Báo cáo hiện giữ delta toàn phiên để tránh đánh đồng các cửa sổ khác nhau.'),
        ]),
        ('7.6 Số kết nối và quy mô mô hình', [
            H('Peering toàn lưới giữa N VPC'),
            E(13, [sub('E', 'P'), ' = ', frac('N(N − 1)', '2')], r'E_P=\frac{N(N-1)}{2}'),
            T('Ký hiệu: N là số VPC (N ≥ 2); E_P là số peering connection cần để mọi cặp nối trực tiếp. Mỗi VPC có N − 1 đối tác, nhưng nhân N(N − 1) đếm mỗi cặp hai lần nên chia 2. Một peering connection hỗ trợ hai chiều; không cần hai connection cho A→B và B→A.'),
            T('Thế số N = 4: E_P = 4 × 3/2 = 6, đúng các cặp AB, AC, AD, BC, BD và CD. Khi thêm VPC thứ năm cần thêm bốn peering, đưa tổng lên 10.'),
            H('Một TGW và attachment của từng VPC'),
            E(14, [sub('E', 'T'), ' = N'], r'E_T=N'),
            T('Ký hiệu: E_T là số VPC attachment trong mô hình mỗi VPC gắn một lần vào cùng một TGW. N = 4 nên E_T = 4, cùng một TGW. Khi thêm VPC thứ năm, thêm một attachment. Đây là đếm tài nguyên kết nối logic, không phải số luồng TCP hoặc số liên kết vật lý của AWS.'),
            H('Số route liên VPC theo cấu hình của đồ án'),
            E(15, [sub('R', 'VPC'), ' = N(N − 1)'], r'R_{\mathrm{VPC}}=N(N-1)'),
            T('Ký hiệu: R_VPC là tổng số entry route có hướng trong các route table của VPC. Giả định mỗi VPC có một route table dùng cho subnet đo và mỗi CIDR VPC đích có một entry riêng. Với N = 4, mỗi bảng có ba route tới các VPC khác, tổng 12 entry ở cả hai mode.'),
            T('Không tính route local, route Internet và entry trong route table TGW vào R_VPC. Nếu tổng hợp CIDR, dùng nhiều route table hoặc khác thiết kế thì số entry có thể khác. TGW giảm số kết nối trong mô hình này; không tự động làm toàn bộ số route cấu hình giảm từ 12 xuống 4.'),
            ('table', ['Số VPC N','Peering E_P','Attachment E_T','Route R_VPC'], [['2','1','2','2'],['4','6','4','12'],['5','10','5','20'],['10','45','10','90']], [1.4,1.7,1.8,1.8]),
            T('Với N tăng, số peering toàn lưới tăng bậc hai, còn số attachment tăng tuyến tính. Hai phương án vẫn có thể cho cùng mọi cặp VPC liên lạc và chạy cùng tập luồng; số tài nguyên khác nhau không có nghĩa phép đo phải dùng số nguồn khác nhau. Đếm kết nối không dự đoán được RTT hoặc thông lượng.'),
        ]),
        ('7.7 Mô hình vật lý giải thích độ trễ', [
            T('Các công thức dưới đây mô tả liên kết và đường đi lý tưởng để giải thích thành phần độ trễ [4]. Chúng chưa được thế bằng tham số vật lý của AWS; hai ví dụ số đều là giả định minh họa.'),
            E(16, [sub('d', 'tx'), ' = ', frac('L', 'R')], r'd_{\mathrm{tx}}=\frac{L}{R}'),
            T('Ký hiệu: d_tx là thời gian đẩy toàn bộ gói lên một liên kết, đơn vị s; L là độ dài gói ở lớp đang xét, đơn vị bit; R là tốc độ truyền của liên kết đó, đơn vị bit/s. Ví dụ giả định L = 1500 × 8 = 12 000 bit và R = 10 × 10⁹ bit/s: d_tx = 1,2 × 10⁻⁶ s = 1,2 µs.'),
            T('MTU 1500 không có nghĩa mọi gói đều dài 1500 byte. Ping dùng payload 56 byte còn có header. Khi tính ở lớp vật lý phải xác định cả overhead thích hợp. Không thay R bằng goodput iperf3 đã đo: goodput và tốc độ của một liên kết là hai đại lượng khác nhau.'),
            E(17, [sub('d', 'prop'), ' = ', frac('ℓ', 'v')], r'd_{\mathrm{prop}}=\frac{\ell}{v}'),
            T('Ký hiệu: d_prop là thời gian tín hiệu lan truyền trên một liên kết, đơn vị s; ℓ là chiều dài đường truyền, đơn vị m; v là tốc độ lan truyền tín hiệu trong môi trường, đơn vị m/s. Ví dụ giả định ℓ = 1000 m, v = 2 × 10⁸ m/s: d_prop = 5 × 10⁻⁶ s = 5 µs. Đây không phải khoảng cách hoặc tốc độ AWS đã đo.'),
            E(18, ['RTT ≈ ', sub('d', 'A→B'), ' + ', sub('d', 'B→A'), ' + ', sub('d', 'end')], r'RTT\approx d_{A\to B}+d_{B\to A}+d_{\mathrm{end}}'),
            T('Ký hiệu: d_A→B và d_B→A là tổng độ trễ đường đi ở từng chiều, đơn vị ms nếu so với ping; mỗi tổng gồm truyền gói, lan truyền, xử lý trung gian và chờ hàng đợi trên đường tương ứng. d_end gom thời gian xử lý tại endpoint và phần đo thời gian ngoài hai đường đi. Hai chiều không nhất thiết có độ trễ bằng nhau.'),
            T('Áp dụng cho đồ án: chênh lệch RTT 0,3552 ms là chênh lệch end-to-end quan sát được. Nó có thể liên quan đường định tuyến và nhiều thành phần khác, chưa thể coi toàn bộ là thời gian xử lý riêng của TGW. Không có chiều dài đường đi, tốc độ từng liên kết hoặc dữ liệu hàng đợi nội bộ nên chưa tách được từng thành phần.'),
            T('Chưa áp dụng mô hình hàng đợi M/M/1 hoặc định luật Little để suy ra thời gian chờ: bộ dữ liệu chưa có số gói trong hàng đợi, tốc độ đến/phục vụ tại nút và điều kiện ổn định cần thiết. Việc thêm công thức chỉ có giá trị khi giả định và dữ liệu đầu vào được kiểm chứng.'),
            R('[4] MIT OpenCourseWare. 6.02 Lecture 17 Packet switching, 2012. Phần transmission và propagation delay.', MIT_URL),
        ]),
    ]


def add_formula_pages(doc, summary, paragraph, table, next_page, hyperlink):
    for title, blocks in pages(summary):
        next_page(doc, title)
        for block in blocks:
            if block[0] == 'text': paragraph(doc, block[1])
            elif block[0] == 'heading': doc.add_heading(block[1], 2)
            elif block[0] == 'equation': equation(doc, block[2], block[1])
            elif block[0] == 'table': table(doc, *block[1:])
            elif block[0] == 'reference':
                p = paragraph(doc, ''); hyperlink(p, block[1], block[2]); p.add_run(' Truy cập ngày 06/10/2026.')


def export_markdown(summary, root_path):
    lines = ['# Công thức và ký hiệu dùng trong phân tích NT531', '',
             'Tài liệu này đi kèm mục 7 của chương kết quả Word. Các phép tính sử dụng dữ liệu thực nghiệm đã kiểm tra; giữ nguyên dữ liệu gốc.', '']
    for title, blocks in pages(summary):
        lines += ['## ' + title, '']
        for block in blocks:
            kind = block[0]
            if kind == 'text': lines += [block[1], '']
            elif kind == 'heading': lines += ['### ' + block[1], '']
            elif kind == 'equation': lines += ['$$', block[3] + rf'\qquad({block[1]})', '$$', '']
            elif kind == 'reference': lines += [f'[{block[1]}]({block[2]}) Truy cập ngày 06/10/2026.', '']
            elif kind == 'table':
                lines += ['| ' + ' | '.join(block[1]) + ' |', '| ' + ' | '.join(['---'] * len(block[1])) + ' |']
                lines += ['| ' + ' | '.join(row) + ' |' for row in block[2]]
                lines += ['']
    lines += ['Định nghĩa bộ đếm ENA [1]: [AWS Monitor network performance for ENA](https://docs.aws.amazon.com/AWSEC2/latest/UserGuide/monitoring-network-performance-ena.html).', '',
              'Nguồn tính toán đầy đủ: `results/analysis/report-metrics.json`, mục `quantitative_methods`. Tái tạo bằng `scripts/build-results-report.py` và `scripts/report_formulas.py`.', '']
    (root_path / 'docs/14-cong-thuc-va-ky-hieu.md').write_text('\n'.join(lines), encoding='utf-8')
